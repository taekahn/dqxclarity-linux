"""Coverage for the community import download / parse edge paths (translate/community.py).

The existing tests cover parse_translation_json shapes (test_community_json.py) and the GAP #15/#18
regressions (test_data_gaps.py). This file fills the remaining thin edges:

* import_custom_zip globs ONLY ``*/json/*.json`` members — directory entries, non-json members, and
  json files outside a json/ directory are ignored — and imports the usable rows;
* import_translation_tarball streams a small in-memory tar.gz, importing only ``json/_lang/en/*.json``
  members, skipping unparsable/empty files, and BATCHING store_many at the batch boundary;
* parse_translation_json handles BOTH the nested ``{id:{ja:en}}`` and the flat ``{ja:en}`` shapes
  (including a mixed file) and applies _is_untranslated to each;
* _is_untranslated drops en==ja, empties, and the placeholder markers, case/space-insensitively;
* parse_merge_xlsx column detection (fixed-english header variants) and the "Story So Far" DeepL
  fallback — including the echo-back edge where DeepL just repeats the JA untranslated.

Every download is mocked (in-memory zip/tar/xlsx bytes via monkeypatched ``community.httpx.get``,
matching the pattern in test_data_gaps.py / test_community_json.py). Nothing here touches the network
and all caches use tmp_path.
"""

from __future__ import annotations

import io
import json
import tarfile
import zipfile

import openpyxl

from dqxclarity.translate import community
from dqxclarity.translate.community import (
    _is_untranslated,
    import_custom_zip,
    import_translation_tarball,
    parse_merge_xlsx,
    parse_translation_json,
)
from dqxclarity.translate.db import TranslationCache


def _b(obj) -> bytes:
    return json.dumps(obj, ensure_ascii=False).encode("utf-8")


def _make_zip(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, content in files.items():
            zf.writestr(name, content)
    return buf.getvalue()


def _make_tarball(files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w:gz") as tar:
        for name, content in files.items():
            info = tarfile.TarInfo(name=name)
            info.size = len(content)
            tar.addfile(info, io.BytesIO(content))
    return buf.getvalue()


def _make_merge_xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------------------------------
# import_custom_zip — glob only */json/*.json, ignore everything else
# --------------------------------------------------------------------------------------------------


def test_import_custom_zip_globs_only_json_members(tmp_path):
    # Only members matching `/json/.+\.json$` are imported. A directory entry, a csv member, and a
    # json file that is NOT under a json/ directory must all be ignored (community.py:38, 239).
    content = _make_zip(
        {
            "repo-main/json/a.json": _b({"看板": "Sign"}),
            "repo-main/json/sub/b.json": _b({"村人": "Villager"}),  # nested under json/ -> imported
            "repo-main/json/": b"",  # directory entry -> ignored
            "repo-main/csv/merge.xlsx": b"not json",  # not json -> ignored
            "repo-main/README.json": _b({"x": "y"}),  # json but outside json/ -> ignored
            "repo-main/json/notes.txt": b"hi",  # under json/ but not .json -> ignored
        }
    )
    cache = TranslationCache(tmp_path / "c.db")
    files, rows = import_custom_zip(content, cache)
    assert files == 2 and rows == 2
    assert cache.lookup("看板") == "Sign"
    assert cache.lookup("村人") == "Villager"
    assert cache.lookup("x") is None  # README.json (outside json/) not imported
    cache.close()


def test_import_custom_zip_skips_unparsable_member(tmp_path):
    # A json member that fails to parse contributes 0 rows and is skipped (parse_translation_json
    # never raises); the file count only counts members that yielded rows.
    content = _make_zip(
        {
            "repo-main/json/good.json": _b({"鍵": "Key"}),
            "repo-main/json/broken.json": b"{not valid json",  # parse -> [] -> skipped
            "repo-main/json/empty.json": _b({"あ": ""}),  # all entries untranslated -> [] -> skipped
        }
    )
    cache = TranslationCache(tmp_path / "c.db")
    files, rows = import_custom_zip(content, cache)
    assert files == 1 and rows == 1
    assert cache.lookup("鍵") == "Key"
    cache.close()


def test_import_custom_zip_empty_archive_no_store(tmp_path):
    # An archive with no usable json members imports nothing and never calls store_many with an
    # empty buffer.
    calls: list[int] = []

    class _CountCache(TranslationCache):
        def store_many(self, rows):
            calls.append(len(rows))
            super().store_many(rows)

    content = _make_zip({"repo-main/README.md": b"hi"})
    cache = _CountCache(tmp_path / "c.db")
    files, rows = import_custom_zip(content, cache)
    assert files == 0 and rows == 0
    assert calls == []  # no store_many on an empty buffer
    cache.close()


# --------------------------------------------------------------------------------------------------
# import_translation_tarball — stream gzip, filter to en/, skip junk, batch store_many
# --------------------------------------------------------------------------------------------------


def test_import_translation_tarball_filters_and_skips(tmp_path):
    # Only json/_lang/en/*.json members import; other langs, non-json, and an unparsable en file are
    # skipped. The unparsable file contributes 0 rows and is NOT counted as an imported file.
    tb = _make_tarball(
        {
            "r-main/json/_lang/en/quests.json": _b({"1": {"新種発見？": "New Species?"}}),
            "r-main/json/_lang/en/items.json": _b({"スライム": "Slime", "あ": ""}),  # 1 usable
            "r-main/json/_lang/en/broken.json": b"{bad",  # parse -> [] -> skipped, 0 files
            "r-main/json/_lang/de/foo.json": _b({"x": "y"}),  # wrong lang -> ignored
            "r-main/json/_lang/en/sub/deep.json": _b({"z": "Z"}),  # nested under en/ but not
            # directly in en/ -> the [^/]+ in _EN_JSON_RE excludes a deeper path
            "r-main/README.md": b"readme",  # non-json -> ignored
        }
    )
    cache = TranslationCache(tmp_path / "c.db")
    files, rows = import_translation_tarball(tb, cache)
    # quests.json (1) + items.json (1 usable) = 2 files, 2 rows. broken/de/deep/readme excluded.
    assert files == 2 and rows == 2
    assert cache.lookup("新種発見？") == "New Species?"
    assert cache.lookup("スライム") == "Slime"
    assert cache.lookup("x") is None  # de/ not imported
    assert cache.lookup("z") is None  # deeper-than-en/ path not matched by _EN_JSON_RE
    cache.close()


def test_import_translation_tarball_batches_store_many(tmp_path):
    # With batch=2 and three single-row files, store_many flushes once at the boundary (the buffer
    # reaches >= batch) and once at the end — proving the streaming batch logic (community.py:210-214).
    calls: list[int] = []

    class _CountCache(TranslationCache):
        def store_many(self, rows):
            calls.append(len(rows))
            super().store_many(rows)

    tb = _make_tarball(
        {
            "r-main/json/_lang/en/a.json": _b({"あ": "A"}),
            "r-main/json/_lang/en/b.json": _b({"い": "I"}),
            "r-main/json/_lang/en/c.json": _b({"う": "U"}),
        }
    )
    cache = _CountCache(tmp_path / "c.db")
    files, rows = import_translation_tarball(tb, cache, batch=2)
    assert files == 3 and rows == 3
    # First flush at the batch boundary (2 buffered), final flush of the remaining 1.
    assert calls == [2, 1]
    assert cache.lookup("あ") == "A"
    assert cache.lookup("う") == "U"
    cache.close()


def test_import_translation_tarball_no_en_files_no_store(tmp_path):
    # A tarball with no en/ members imports nothing and does not call store_many on an empty buffer.
    calls: list[int] = []

    class _CountCache(TranslationCache):
        def store_many(self, rows):
            calls.append(len(rows))
            super().store_many(rows)

    tb = _make_tarball(
        {
            "r-main/json/_lang/de/foo.json": _b({"x": "y"}),
            "r-main/README.md": b"hi",
        }
    )
    cache = _CountCache(tmp_path / "c.db")
    files, rows = import_translation_tarball(tb, cache)
    assert files == 0 and rows == 0
    assert calls == []
    cache.close()


# --------------------------------------------------------------------------------------------------
# parse_translation_json — nested + flat in one file, with skip rules
# --------------------------------------------------------------------------------------------------


def test_parse_translation_json_mixed_nested_and_flat_in_one_file():
    # A single file with BOTH shapes: a nested {id:{ja:en}} entry and flat {ja:en} entries. Both are
    # parsed and _is_untranslated is applied to each (the identical and empty entries drop out).
    data = _b(
        {
            "10001": {"新種発見？": "New Species?"},  # nested -> kept
            "スライム": "Slime",  # flat -> kept
            "同じ": "同じ",  # flat en==ja -> dropped
            "空": "",  # flat empty -> dropped
            "10002": {"霊魂": "霊魂"},  # nested en==ja -> dropped
        }
    )
    rows = dict(parse_translation_json(data))
    assert rows == {"新種発見？": "New Species?", "スライム": "Slime"}


def test_parse_translation_json_multi_entry_nested_dict():
    # A nested value may hold MORE than one ja->en pair; every usable inner pair is taken
    # (community.py:180-184), and untranslated inner pairs are individually dropped.
    data = _b({"grp": {"剣": "Sword", "盾": "Shield", "槍": ""}})
    rows = dict(parse_translation_json(data))
    assert rows == {"剣": "Sword", "盾": "Shield"}


# --------------------------------------------------------------------------------------------------
# _is_untranslated — direct unit coverage of the skip rules
# --------------------------------------------------------------------------------------------------


def test_is_untranslated_drops_empty_identical_and_placeholders():
    assert _is_untranslated("x", "") is True  # empty EN
    assert _is_untranslated("同じ", "同じ") is True  # EN identical to JA
    assert _is_untranslated("x", "untranslated") is True
    assert _is_untranslated("x", "No Translation") is True  # case-insensitive
    assert _is_untranslated("x", "  NULL  ") is True  # whitespace-padded placeholder
    assert _is_untranslated("x", "none") is True


def test_is_untranslated_keeps_real_translations():
    assert _is_untranslated("スライム", "Slime") is False
    # A real EN that merely CONTAINS a placeholder word (not equal to it) is kept.
    assert _is_untranslated("x", "A null pointer appeared") is False


# --------------------------------------------------------------------------------------------------
# parse_merge_xlsx — column detection + Story So Far DeepL fallback (incl. echo-back edge)
# --------------------------------------------------------------------------------------------------


def test_parse_merge_xlsx_detects_english_translation_header_variant():
    # The fixed-english column is found by a header containing "english" AND ("fixed" OR
    # "translation") (community.py:54-58). A bare "English" header (neither word) is NOT a match, so
    # such a sheet yields no rows.
    data = _make_merge_xlsx(
        {
            "Dialogue": [
                ["Japanese", "English Translation"],  # "english"+"translation" -> col 1
                ["会話", "Dialogue line"],
            ],
            "Quests": [
                ["Japanese", "English"],  # bare "English" -> no column -> sheet skipped
                ["クエスト", "Quest line"],
            ],
        }
    )
    rows = dict(parse_merge_xlsx(data))
    assert rows == {"会話": "Dialogue line"}


def test_parse_merge_xlsx_skips_sheet_without_english_column():
    # A sheet whose header has no usable english column is skipped entirely (en_col is None).
    data = _make_merge_xlsx(
        {
            "Dialogue": [
                ["Japanese", "Notes"],  # no english col -> skipped
                ["会話", "irrelevant"],
            ],
        }
    )
    assert parse_merge_xlsx(data) == []


def test_parse_merge_xlsx_ignores_unknown_sheets():
    # Only the four known sheets are read; an unrelated sheet is never parsed even if it has an
    # english column (community.py:51, 95-96).
    data = _make_merge_xlsx(
        {
            "Glossary": [  # not in _SHEETS
                ["Japanese", "Fixed English"],
                ["用語", "Term"],
            ],
            "Dialogue": [
                ["Japanese", "Fixed English"],
                ["会話", "Dialogue"],
            ],
        }
    )
    rows = dict(parse_merge_xlsx(data))
    assert rows == {"会話": "Dialogue"}


def test_parse_merge_xlsx_story_so_far_deepl_fallback_without_deepl_header():
    # "Story So Far" falls back to the column immediately BEFORE fixed-english when no header
    # contains "deepl" (community.py:62-72). Here the MT column is headed "MT", not "DeepL", so the
    # fallback is positional (en_col - 1).
    data = _make_merge_xlsx(
        {
            "Story So Far": [
                ["Source", "MT", "Fixed English"],  # no "deepl" header -> positional fallback
                ["物語1", "MT Recap", None],  # fixed-en empty -> use col 1 (MT)
                ["物語2", "MT Recap 2", "Fixed Recap"],  # fixed-en present -> preferred
            ],
        }
    )
    rows = dict(parse_merge_xlsx(data))
    assert rows["物語1"] == "MT Recap"
    assert rows["物語2"] == "Fixed Recap"
    assert len(rows) == 2


def test_parse_merge_xlsx_story_so_far_drops_echo_back_deepl():
    # The Story So Far DeepL fallback must NOT import an untranslated cell: when DeepL just echoes the
    # JA back (deepl == ja), _is_untranslated drops it so no untranslated recap is imported
    # (community.py:110-115).
    data = _make_merge_xlsx(
        {
            "Story So Far": [
                ["Source (Japanese)", "DeepL Translation", "Fixed English"],
                ["物語echo", "物語echo", None],  # DeepL == JA -> untranslated -> dropped
                ["物語ok", "DeepL OK", None],  # usable DeepL -> kept
                ["物語null", None, None],  # neither -> dropped
            ],
        }
    )
    rows = dict(parse_merge_xlsx(data))
    assert "物語echo" not in rows
    assert "物語null" not in rows
    assert rows["物語ok"] == "DeepL OK"
    assert len(rows) == 1


# --------------------------------------------------------------------------------------------------
# sync wrappers — single mocked fetch drives the in-memory parsers (no network)
# --------------------------------------------------------------------------------------------------


def test_sync_all_static_streams_tarball(tmp_path, monkeypatch):
    # sync_all_static fetches the whole dqx_translations tarball ONCE (mocked) and imports its entire
    # en/ corpus via import_translation_tarball.
    class _Resp:
        def __init__(self, content: bytes):
            self.content = content

        def raise_for_status(self):
            pass

    tb = _make_tarball(
        {
            "r-main/json/_lang/en/a.json": _b({"スライム": "Slime"}),
            "r-main/json/_lang/en/b.json": _b({"1": {"勇者": "Hero"}}),
            "r-main/json/_lang/fr/c.json": _b({"x": "y"}),  # ignored
        }
    )
    seen: list[str] = []

    def fake_get(url, **kwargs):
        seen.append(url)
        return _Resp(tb)

    monkeypatch.setattr(community.httpx, "get", fake_get)
    cache = TranslationCache(tmp_path / "c.db")
    files, rows = community.sync_all_static(cache)
    assert seen == [community.TRANSLATIONS_TARBALL_URL]
    assert files == 2 and rows == 2
    assert cache.lookup("スライム") == "Slime"
    assert cache.lookup("勇者") == "Hero"
    assert cache.lookup("x") is None
    cache.close()


def test_sync_community_imports_merge_xlsx(tmp_path, monkeypatch):
    # sync_community fetches merge.xlsx ONCE (mocked) and imports the parsed rows at the community
    # tier; returns the row count.
    class _Resp:
        def __init__(self, content: bytes):
            self.content = content

        def raise_for_status(self):
            pass

    data = _make_merge_xlsx(
        {
            "Dialogue": [
                ["Japanese", "Fixed English"],
                ["会話1", "Line one"],
                ["会話2", "Line two"],
            ],
        }
    )
    seen: list[str] = []

    def fake_get(url, **kwargs):
        seen.append(url)
        return _Resp(data)

    monkeypatch.setattr(community.httpx, "get", fake_get)
    cache = TranslationCache(tmp_path / "c.db")
    count = community.sync_community(cache)
    assert seen == [community.MERGE_XLSX_URL]
    assert count == 2
    assert cache.lookup("会話1") == "Line one"
    assert cache.source_of("会話1") == "community"  # imported at the community quality tier
    cache.close()
